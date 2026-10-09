"""Render paper.md into LaTeX, Word and reading-copy Markdown.

    python build.py

paper.md is the only place the prose lives. Figure, table and equation
numbers are assigned here from the order the directive blocks appear, and
every cross-reference is resolved against them, so a block cannot end up
numbered one way in the PDF and another way in the Word file. Citations are
expanded from references.py for the same reason.

The build fails loudly on an unresolved cross-reference or an unknown
citation key rather than emitting a document with a question mark in it.

Outputs, all beside this file:

    main.tex            single compile on Overleaf, no bibtex pass needed
    refs.bib            for anyone who would rather run bibtex
    NMAFC-paper.docx    Word, with display equations rendered as images
    paper-rendered.md   reading copy with citations and numbers expanded
"""

from __future__ import annotations

import pathlib
import re
import sys

from references import REFERENCES

HERE = pathlib.Path(__file__).parent
SOURCE = HERE / "paper.md"
EQ_DIR = HERE / "figures" / "equations"

# LaTeX characters that mean something other than themselves. Applied only
# to prose, never inside maths or a code span.
TEX_ESCAPES = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$",
    "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
    "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}

# Enough of the maths alphabet to render inline maths as plain text in the
# Word file, which cannot typeset it. Display maths goes through matplotlib
# instead and keeps its notation.
UNICODE_MATH = {
    r"\alpha": "α", r"\beta": "β", r"\gamma": "γ",
    r"\delta": "δ", r"\eta": "η", r"\theta": "θ",
    r"\lambda": "λ", r"\tau": "τ", r"\omega": "ω",
    r"\Delta": "Δ", r"\ell": "ℓ", r"\in": "∈",
    r"\le": "≤", r"\ge": "≥", r"\approx": "≈",
    r"\cdot": "·", r"\times": "×", r"\sqrt": "√",
    r"\leftarrow": "←", r"\rightarrow": "→", r"\gets": "←",
    r"\min": "min", r"\mathrm": "", r"\mathcal": "", r"\,": " ",
    r"\;": " ",
}
# Unicode has no subscript for every letter, so anything without one keeps
# the underscore rather than silently losing the level.
SUBSCRIPTS = {"0": "₀", "1": "₁", "2": "₂", "3": "₃", "4": "₄",
              "5": "₅", "6": "₆", "7": "₇", "8": "₈", "9": "₉",
              "h": "ₕ", "i": "ᵢ", "j": "ⱼ", "k": "ₖ", "l": "ₗ",
              "m": "ₘ", "n": "ₙ", "p": "ₚ", "r": "ᵣ", "s": "ₛ",
              "t": "ₜ", "x": "ₓ"}


# --------------------------------------------------------------- parsing

class Block:
    """One directive block: a figure, a table or a display equation."""

    def __init__(self, kind: str, ident: str, attrs: dict,
                 body: list[str]) -> None:
        self.kind = kind
        self.id = ident
        self.attrs = attrs
        self.body = body
        self.number = 0


def parse(text: str) -> tuple[dict, list]:
    """Split the source into a metadata dict and a list of items.

    An item is either a string of Markdown or a Block.
    """
    meta: dict[str, str] = {}
    items: list = []
    lines = text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith("::: "):
            header = line[4:].strip()
            body: list[str] = []
            i += 1
            while i < len(lines) and lines[i].strip() != ":::":
                body.append(lines[i])
                i += 1
            i += 1
            parts = header.split()
            kind = parts[0]
            attrs = dict(p.split("=", 1) for p in parts[1:] if "=" in p)
            if kind == "meta":
                for entry in body:
                    if ":" in entry:
                        key, value = entry.split(":", 1)
                        meta[key.strip()] = value.strip()
                continue
            items.append(Block(kind, attrs.get("id", ""), attrs, body))
            continue
        if line.startswith("# ") and "title" not in meta:
            meta["title"] = line[2:].strip()
        items.append(line)
        i += 1
    return meta, items


def number_blocks(items: list) -> dict[str, tuple[str, int]]:
    """Assign numbers in order of appearance and build the label map."""
    counters = {"figure": 0, "table": 0, "equation": 0}
    prefix = {"figure": "fig", "table": "tab", "equation": "eq"}
    labels: dict[str, tuple[str, int]] = {}
    for item in items:
        if isinstance(item, Block):
            counters[item.kind] += 1
            item.number = counters[item.kind]
            labels[f"{prefix[item.kind]}:{item.id}"] = (item.kind,
                                                        item.number)
    return labels


# ------------------------------------------------------------- citations

def author_short(entry: dict) -> str:
    """Harvard in-text author string: one, two, or the first plus et al."""
    families = [family for family, _ in entry["authors"]]
    if entry.get("etal") or len(families) > 2:
        return f"{families[0]} et al."
    if len(families) == 2:
        return f"{families[0]} and {families[1]}"
    return families[0]


def cite_keys(raw: str) -> list[str]:
    return [k.strip().lstrip("@") for k in raw.split(";") if k.strip()]


def expand_citations(text: str, used: set[str]) -> str:
    """Turn [@key] into (Author, year) and [-@key] into Author (year)."""

    def narrative(match: re.Match) -> str:
        key = match.group(1)
        used.add(key)
        return f"{author_short(REFERENCES[key])} ({REFERENCES[key]['year']})"

    def parenthetical(match: re.Match) -> str:
        rendered = []
        for key in cite_keys(match.group(1)):
            used.add(key)
            entry = REFERENCES[key]
            rendered.append(f"{author_short(entry)}, {entry['year']}")
        return "(" + "; ".join(rendered) + ")"

    text = re.sub(r"\[-@([A-Za-z0-9]+)\]", narrative, text)
    return re.sub(r"\[@([^\]]+)\]", parenthetical, text)


def latex_citations(text: str, used: set[str]) -> str:
    """The same, as natbib commands."""

    def narrative(match: re.Match) -> str:
        used.add(match.group(1))
        return "\\citet{" + match.group(1) + "}"

    def parenthetical(match: re.Match) -> str:
        keys = cite_keys(match.group(1))
        used.update(keys)
        return "\\citep{" + ",".join(keys) + "}"

    text = re.sub(r"\[-@([A-Za-z0-9]+)\]", narrative, text)
    return re.sub(r"\[@([^\]]+)\]", parenthetical, text)


def check_keys(text: str) -> None:
    found = set()
    for match in re.finditer(r"\[-?@([^\]]+)\]", text):
        found.update(cite_keys(match.group(1)))
    unknown = sorted(found - set(REFERENCES))
    if unknown:
        sys.exit(f"unknown citation keys: {', '.join(unknown)}")


# ------------------------------------------------------- cross-references

def expand_refs(text: str, labels: dict, latex: bool) -> str:
    """Resolve {@fig:x} to a number, or to \\ref for the LaTeX build."""

    def one(match: re.Match) -> str:
        label = match.group(1)
        if label not in labels:
            sys.exit(f"unresolved cross-reference: {{@{label}}}")
        kind, number = labels[label]
        if latex:
            return "\\ref{" + kind + ":" + label.split(":", 1)[1] + "}"
        return str(number)

    return re.sub(r"\{@([a-z]+:[A-Za-z0-9_-]+)\}", one, text)


# ------------------------------------------------------------ references

def render_harvard(key: str) -> str:
    """One reference-list entry, Harvard, British spelling."""
    e = REFERENCES[key]
    names = []
    for family, given in e["authors"]:
        initials = "".join(f"{p[0]}." for p in given.replace("-", " ").split()
                           if p)
        names.append(f"{family}, {initials}")
    if e.get("etal"):
        names.append("et al.")
    if len(names) == 1:
        authors = names[0]
    else:
        authors = ", ".join(names[:-1]) + " and " + names[-1]
    if names[-1] == "et al.":
        authors = ", ".join(names[:-1]) + " " + names[-1]
    out = [f"{authors} ({e['year']}) '{e['title']}'"]
    if e["type"] == "article":
        out.append(f", {e['journal']}")
        if e.get("volume"):
            out.append(f", {e['volume']}")
        if e.get("number"):
            out.append(f"({e['number']})")
        if e.get("pages"):
            out.append(f", pp. {e['pages'].replace('--', '-')}")
        out.append(".")
    elif e["type"] in ("inproceedings", "incollection"):
        out.append(f", in {e['booktitle']}")
        if e.get("volume") and e["type"] == "inproceedings":
            out.append(f", {e['volume']}")
        if e.get("pages"):
            out.append(f", pp. {e['pages'].replace('--', '-')}")
        if e.get("publisher"):
            out.append(f". {e['publisher']}")
        out.append(".")
    elif e["type"] == "preprint":
        out.append(f" [Preprint]. {e['archive']}:{e['eprint']}.")
        out.append(f" Posted {e['posted']}.")
    if e.get("doi"):
        out.append(f" doi: {e['doi']}.")
    elif e.get("eprint") and e["type"] != "preprint":
        out.append(f" arXiv:{e['eprint']}.")
    if e["type"] == "preprint":
        out.append(f" Available at: https://arxiv.org/abs/{e['eprint']}")
        out.append(" (Accessed: 10 September 2026).")
    return "".join(out)


def sorted_keys(used: set[str]) -> list[str]:
    def key_of(k: str) -> tuple:
        e = REFERENCES[k]
        return (e["authors"][0][0].lower(), e["year"], e["title"].lower())
    return sorted(used, key=key_of)


def render_bibtex(key: str) -> str:
    e = REFERENCES[key]
    kind = {"preprint": "misc"}.get(e["type"], e["type"])
    authors = " and ".join(f"{family}, {given}"
                           for family, given in e["authors"])
    if e.get("etal"):
        authors += " and others"
    fields = [("author", authors), ("title", e["title"]),
              ("year", str(e["year"]))]
    for name in ("journal", "booktitle", "volume", "number", "pages",
                 "publisher", "doi", "eprint"):
        if e.get(name):
            fields.append((name, str(e[name])))
    if e["type"] == "preprint":
        fields.append(("note", f"{e['archive']} preprint, "
                               f"posted {e['posted']}"))
        fields.append(("archivePrefix", "arXiv"))
    body = ",\n".join(f"  {n} = {{{v}}}" for n, v in fields)
    return f"@{kind}{{{key},\n{body}\n}}\n"


# ----------------------------------------------------------- inline text

SPANS = re.compile(r"`([^`]+)`"
                   r"|\$([^$]+)\$"
                   r"|(\[-?@[^\]]+\]|\{@[a-z]+:[A-Za-z0-9_-]+\})")


def split_spans(text: str) -> list[tuple[str, str]]:
    """Cut a line into text, maths, code and marker spans.

    A marker is an unexpanded citation or cross-reference. It is carried
    through the escaping step untouched so that the expansion which follows
    still recognises it.
    """
    spans: list[tuple[str, str]] = []
    position = 0
    for match in SPANS.finditer(text):
        if match.start() > position:
            spans.append(("text", text[position:match.start()]))
        if match.group(1) is not None:
            spans.append(("code", match.group(1)))
        elif match.group(2) is not None:
            spans.append(("math", match.group(2)))
        else:
            spans.append(("marker", match.group(3)))
        position = match.end()
    if position < len(text):
        spans.append(("text", text[position:]))
    return spans


def tex_escape(text: str) -> str:
    return "".join(TEX_ESCAPES.get(char, char) for char in text)


def to_latex_inline(text: str) -> str:
    out = []
    for kind, content in split_spans(text):
        if kind == "code":
            out.append("\\texttt{" + tex_escape(content) + "}")
        elif kind in ("math", "marker"):
            out.append(content if kind == "marker"
                       else "$" + content + "$")
        else:
            escaped = tex_escape(content)
            escaped = escaped.replace('"', "''")
            out.append(escaped)
    return "".join(out)


def to_plain_inline(text: str) -> str:
    """Inline text for Word: maths flattened to Unicode, code left bare."""
    out = []
    for kind, content in split_spans(text):
        if kind == "math":
            flat = content
            for token, replacement in sorted(UNICODE_MATH.items(),
                                             key=lambda kv: -len(kv[0])):
                flat = flat.replace(token, replacement)
            flat = re.sub(r"_\{?([A-Za-z0-9]+)\}?",
                          lambda m: SUBSCRIPTS.get(m.group(1),
                                                   "_" + m.group(1)), flat)
            flat = flat.replace("{", "").replace("}", "").strip()
            out.append(re.sub(r"(?<=[ΔΣΠ]) (?=[a-zA-Z])", "", flat))
        else:
            out.append(content)
    return "".join(out)


# ------------------------------------------------------------ table bits

def parse_table(body: list[str]) -> tuple[str, list[list[str]]]:
    """Split a table block into its caption and its rows."""
    caption_lines: list[str] = []
    rows: list[list[str]] = []
    for line in body:
        stripped = line.strip()
        if stripped.startswith("|"):
            cells = [c.strip() for c in stripped.strip("|").split("|")]
            if all(set(c) <= set("-: ") and c for c in cells):
                continue
            rows.append(cells)
        elif stripped:
            caption_lines.append(stripped)
    return " ".join(caption_lines), rows


NUMERIC = re.compile(r"^[+-]?[\d.,]+(e-?\d+)?%?$|^[+-]?[\d.]+ ?\[.*\]$"
                     r"|^[\d,]+ vs [\d,]+$|^[+-]?\d+/\d+$|^n\.s\.$|^-$")

# Tables too wide for a single column, which therefore float across both.
WIDE_TABLES = {"main", "goldlen", "advloss", "temwin", "files", "neg"}


def column_spec(rows: list[list[str]], wide: bool) -> str:
    """Right-align columns that hold only numbers, wrap what will not fit.

    A column of the two-column layout holds roughly 45 characters at this
    size and the full width roughly 95, so the threshold for switching to
    wrapped p columns depends on which the table is going into.
    """
    widths = [max(len(r[i]) for r in rows) for i in range(len(rows[0]))]
    if sum(widths) > (92 if wide else 44):
        total = sum(widths)
        return "".join(
            r"p{%.3f\linewidth}" % (0.95 * w / total) for w in widths)
    spec = []
    for i, _ in enumerate(widths):
        body = [r[i] for r in rows[1:]]
        spec.append("r" if body and all(NUMERIC.match(c) for c in body)
                    else "l")
    return "".join(spec)


# ---------------------------------------------------------------- LaTeX

PREAMBLE = r"""\documentclass[10pt,a4paper,twocolumn]{article}
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage[british]{babel}
\usepackage{lmodern}
\usepackage[margin=0.78in,columnsep=0.3in]{geometry}
\usepackage{amsmath,amssymb}
\usepackage{graphicx}
\usepackage{booktabs}
\usepackage{array}
\usepackage{caption}
\usepackage[round,authoryear]{natbib}
\usepackage[hidelinks]{hyperref}
\captionsetup{font=small,labelfont=bf,skip=4pt}
\setlength{\tabcolsep}{3.5pt}
\renewcommand{\arraystretch}{1.1}
\setlength{\parskip}{0pt}
\setlength{\belowcaptionskip}{-4pt}
\bibpunct{(}{)}{;}{a}{,}{,}

%% Twelve of the floats span both columns. LaTeX will only place a spanning
%% float at the top of a later page, so without loosening these counters it
%% queues them up and stops with "too many unprocessed floats".
\usepackage{stfloats}
\setcounter{topnumber}{3}
\setcounter{bottomnumber}{2}
\setcounter{totalnumber}{4}
\setcounter{dbltopnumber}{3}
\renewcommand{\topfraction}{0.92}
\renewcommand{\dbltopfraction}{0.92}
\renewcommand{\bottomfraction}{0.5}
\renewcommand{\textfraction}{0.07}
\renewcommand{\floatpagefraction}{0.75}
\renewcommand{\dblfloatpagefraction}{0.7}

\begin{document}

%% The title block and the abstract span both columns. Building them by hand
%% rather than through \maketitle keeps the file self-contained, which
%% matters because it is compiled by uploading it on its own.
\twocolumn[{%%
\begin{center}
{\LARGE\bfseries %(title)s\par}
\vspace{0.9em}
{\large %(author)s\par}
\vspace{0.2em}
{\normalsize %(affiliation)s \quad \texttt{%(email)s}\par}
\vspace{0.2em}
{\normalsize %(date)s\par}
\end{center}
\vspace{0.6em}
\begin{center}
\begin{minipage}{0.86\textwidth}
\small
{\centering\textsc{Abstract}\par}
\vspace{0.3em}
%(abstract)s
\end{minipage}
\end{center}
\vspace{1.4em}
}]

"""


def emit_latex(meta: dict, items: list, labels: dict) -> str:
    used: set[str] = set()
    out: list[str] = []
    abstract: list[str] = []
    in_abstract = False
    for item in items:
        if isinstance(item, Block):
            out.append(latex_block(item, labels, used))
            continue
        line = item.rstrip()
        if line.startswith("# "):
            continue
        if line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            title = line.lstrip("# ").strip()
            in_abstract = title.lower() == "abstract"
            if in_abstract:
                continue
            body = re.sub(r"^[\d.]+\s+", "", title)
            starred = "" if re.match(r"^\d", title) else "*"
            if title.lower().startswith("appendix"):
                body = title
                starred = "*"
            command = {2: "section", 3: "subsection"}[min(level, 3)]
            out.append("\\%s%s{%s}\n" % (command, starred,
                                         tex_escape(body)))
            continue
        target = abstract if in_abstract else out
        if not line.strip():
            target.append("\n")
            continue
        target.append(latex_line(line, labels, used) + "\n")
    head = PREAMBLE % {
        "title": tex_escape(meta["title"]),
        "author": tex_escape(meta["author"]),
        "affiliation": tex_escape(meta["affiliation"]),
        "email": meta["email"], "date": meta["date"],
        "abstract": "".join(abstract).strip()}
    body = head + "".join(out)
    body = body.replace("\\section*{References}\n",
                        latex_bibliography(used))
    return body + "\n\\end{document}\n"


def latex_line(text: str, labels: dict, used: set[str]) -> str:
    """Escape the prose first, then expand the markers it left alone."""
    escaped = to_latex_inline(text)
    return expand_refs(latex_citations(escaped, used), labels, latex=True)


def latex_bibliography(used: set[str]) -> str:
    lines = ["\\section*{References}\n",
             "\\begin{thebibliography}{99}\n",
             "\\setlength{\\itemsep}{2pt}\n"]
    for key in sorted_keys(used):
        entry = REFERENCES[key]
        label = tex_escape(f"{author_short(entry)}({entry['year']})")
        text = to_latex_inline(render_harvard(key))
        lines.append("\\bibitem[%s]{%s}\n%s\n\n" % (label, key, text))
    lines.append("\\end{thebibliography}\n")
    return "".join(lines)


def latex_block(block: Block, labels: dict, used: set[str]) -> str:
    tag = block.id
    if block.kind == "equation":
        maths = " ".join(line.strip() for line in block.body if line.strip())
        return ("\\begin{equation}\\label{equation:%s}\n%s\n"
                "\\end{equation}\n" % (tag, maths))
    if block.kind == "figure":
        caption = latex_line(
            " ".join(x.strip() for x in block.body if x.strip()),
            labels, used)
        # Every figure is a multi-panel strip at least 1.8 times as wide as
        # it is tall, so all of them span the two columns.
        return ("\\begin{figure*}[tb]\n\\centering\n"
                "\\includegraphics[width=\\textwidth]{%s.pdf}\n"
                "\\caption{%s}\\label{figure:%s}\n\\end{figure*}\n"
                % (block.attrs["src"], caption, tag))
    caption, rows = parse_table(block.body)
    wide = tag in WIDE_TABLES
    star = "*" if wide else ""
    size = "\\scriptsize" if len(rows[0]) > 5 else "\\small"
    lines = ["\\begin{table%s}[tb]\n\\centering\n" % star,
             "\\caption{%s}" % latex_line(caption, labels, used),
             "\\label{table:%s}\n" % tag, size + "\n",
             "\\begin{tabular}{%s}\n\\toprule\n" % column_spec(rows, wide)]
    for index, row in enumerate(rows):
        cells = [latex_line(c, labels, used) for c in row]
        if cells[0].startswith("["):
            cells[0] = "{}" + cells[0]
        lines.append(" & ".join(cells) + " \\\\\n")
        if index == 0:
            lines.append("\\midrule\n")
    lines.append("\\bottomrule\n\\end{tabular}\n\\end{table%s}\n" % star)
    return "".join(lines)


# ------------------------------------------------------------- Markdown

def emit_markdown(meta: dict, items: list, labels: dict) -> str:
    used: set[str] = set()
    out = [f"# {meta['title']}\n",
           f"\n{meta['author']}. {meta['affiliation']}. {meta['email']}.",
           f" {meta['date']}.\n"]
    for item in items:
        if isinstance(item, Block):
            out.append(markdown_block(item, labels, used))
            continue
        if item.startswith("# "):
            continue
        line = expand_refs(expand_citations(item, used), labels, latex=False)
        out.append(line + "\n")
    text = "".join(out)
    bibliography = "\n".join(f"{render_harvard(k)}\n"
                             for k in sorted_keys(used))
    return text.replace("## References\n", "## References\n\n"
                        + bibliography)


def markdown_block(block: Block, labels: dict, used: set[str]) -> str:
    name = {"figure": "Figure", "table": "Table",
            "equation": "Equation"}[block.kind]
    if block.kind == "equation":
        maths = " ".join(line.strip() for line in block.body if line.strip())
        return f"\n$$\n{maths}\n$$\n\n({block.number})\n"
    if block.kind == "figure":
        caption = expand_refs(expand_citations(
            " ".join(l.strip() for l in block.body if l.strip()), used),
            labels, latex=False)
        return (f"\n![{name} {block.number}]({block.attrs['src']}.png)\n\n"
                f"{name} {block.number}. {caption}\n")
    caption, rows = parse_table(block.body)
    caption = expand_refs(expand_citations(caption, used), labels,
                          latex=False)
    lines = [f"\n{name} {block.number}. {caption}\n\n",
             "| " + " | ".join(rows[0]) + " |\n",
             "|" + "|".join(["---"] * len(rows[0])) + "|\n"]
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |\n")
    return "".join(lines)


# ----------------------------------------------------------------- Word

def render_equation_png(maths: str, path: pathlib.Path) -> None:
    """Word cannot typeset LaTeX, so display maths becomes an image."""
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    figure = plt.figure(figsize=(0.01, 0.01))
    figure.text(0, 0, f"${maths}$", fontsize=15)
    figure.savefig(path, dpi=300, bbox_inches="tight",
                   pad_inches=0.06, transparent=False,
                   facecolor="white")
    plt.close(figure)


def png_inches(path: pathlib.Path, dpi: int = 300) -> float:
    """Width of a PNG in inches, read out of its IHDR chunk."""
    raw = path.read_bytes()[16:20]
    return int.from_bytes(raw, "big") / dpi


def docx_paragraph(document, text: str, style=None):
    from docx.shared import Pt

    paragraph = document.add_paragraph(style=style)
    for kind, content in split_spans(text):
        run = paragraph.add_run(to_plain_inline(
            content if kind == "text" else
            ("`" + content + "`" if kind == "code" else "$" + content + "$")))
        if kind == "code":
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
    return paragraph


def emit_docx(meta: dict, items: list, labels: dict, path: pathlib.Path):
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
    from docx.shared import Inches, Pt

    EQ_DIR.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    normal = document.styles["Normal"]
    normal.font.name = "Cambria"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)

    heading = document.add_paragraph()
    heading.alignment = ALIGN.CENTER
    run = heading.add_run(meta["title"])
    run.bold = True
    run.font.size = Pt(16)
    byline = document.add_paragraph()
    byline.alignment = ALIGN.CENTER
    byline.add_run(f"{meta['author']}\n{meta['affiliation']}\n"
                   f"{meta['email']}\n{meta['date']}")

    used: set[str] = set()
    for item in items:
        if isinstance(item, Block):
            docx_block(document, item, labels, used)
            continue
        if item.startswith("# ") or not item.strip():
            continue
        if item.startswith("#"):
            level = len(item) - len(item.lstrip("#"))
            title = item.lstrip("# ").strip()
            document.add_heading(title, level=max(1, level - 1))
            if title.lower() == "references":
                for key in sorted_keys(used):
                    entry = document.add_paragraph(render_harvard(key))
                    entry.paragraph_format.left_indent = Inches(0.4)
                    entry.paragraph_format.first_line_indent = Inches(-0.4)
                    entry.paragraph_format.space_after = Pt(6)
            continue
        text = expand_refs(expand_citations(item, used), labels, latex=False)
        docx_paragraph(document, text).alignment = ALIGN.JUSTIFY
    document.save(path)
    return used


def docx_block(document, block: Block, labels: dict, used: set[str]) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH as ALIGN
    from docx.shared import Inches, Pt

    name = {"figure": "Figure", "table": "Table",
            "equation": "Equation"}[block.kind]
    if block.kind == "equation":
        maths = " ".join(line.strip() for line in block.body if line.strip())
        image = EQ_DIR / f"eq{block.number}.png"
        render_equation_png(maths, image)
        holder = document.add_table(rows=1, cols=2)
        holder.columns[0].width = Inches(5.4)
        holder.columns[1].width = Inches(0.8)
        cell = holder.cell(0, 0).paragraphs[0]
        cell.alignment = ALIGN.CENTER
        cell.add_run().add_picture(
            str(image), width=Inches(min(5.2, png_inches(image))))
        tail = holder.cell(0, 1).paragraphs[0]
        tail.alignment = ALIGN.RIGHT
        tail.add_run(f"({block.number})")
        return
    if block.kind == "figure":
        caption = expand_refs(expand_citations(
            " ".join(l.strip() for l in block.body if l.strip()), used),
            labels, latex=False)
        holder = document.add_paragraph()
        holder.alignment = ALIGN.CENTER
        holder.add_run().add_picture(
            str(HERE / (block.attrs["src"] + ".png")), width=Inches(6.2))
        line = document.add_paragraph()
        line.alignment = ALIGN.CENTER
        label = line.add_run(f"{name} {block.number}. ")
        label.bold = True
        label.font.size = Pt(9.5)
        body = line.add_run(to_plain_inline(caption))
        body.font.size = Pt(9.5)
        return
    caption, rows = parse_table(block.body)
    caption = expand_refs(expand_citations(caption, used), labels,
                          latex=False)
    line = document.add_paragraph()
    label = line.add_run(f"{name} {block.number}. ")
    label.bold = True
    label.font.size = Pt(9.5)
    body = line.add_run(to_plain_inline(caption))
    body.font.size = Pt(9.5)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = table.cell(r, c).paragraphs[0]
            run = cell.add_run(to_plain_inline(text))
            run.font.size = Pt(8)
            run.bold = r == 0
    document.add_paragraph()


# ----------------------------------------------------------------- main

def prose_words(items: list) -> dict:
    """Count the running text only, which is what a length limit means.

    Table cells, captions and the reference list are excluded, since
    counting them puts thousands of words on the total that no venue
    counts and no reader reads as prose.
    """
    counts = {"body": 0, "appendix": 0}
    where = "body"
    for item in items:
        if isinstance(item, Block):
            continue
        if item.startswith("#"):
            title = item.lstrip("# ").strip().lower()
            if title == "references":
                where = "skip"
            elif title.startswith("appendix"):
                where = "appendix"
            continue
        if where != "skip":
            counts[where] += len(re.findall(r"\S+", item))
    return counts


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8")
    check_keys(source)
    meta, items = parse(source)
    labels = number_blocks(items)


    (HERE / "main.tex").write_text(
        emit_latex(meta, items, labels), encoding="utf-8")
    rendered = emit_markdown(meta, items, labels)
    (HERE / "paper-rendered.md").write_text(rendered, encoding="utf-8")

    cited: set[str] = set()
    expand_citations(source, cited)
    (HERE / "refs.bib").write_text(
        "\n".join(render_bibtex(k) for k in sorted_keys(cited)),
        encoding="utf-8")

    docx_written = True
    try:
        emit_docx(meta, items, labels, HERE / "NMAFC-paper.docx")
    except PermissionError:
        docx_written = False

    unused = sorted(set(REFERENCES) - cited)
    words = prose_words(items)
    counts = {k: sum(1 for i in items if isinstance(i, Block)
                     and i.kind == k)
              for k in ("figure", "table", "equation")}
    print(f"prose words  {words['body']} body, "
          f"{words['appendix']} appendix")
    print(f"figures {counts['figure']}  tables {counts['table']}  "
          f"equations {counts['equation']}")
    print(f"references   {len(cited)} cited of {len(REFERENCES)}")
    if unused:
        print(f"NOT CITED    {', '.join(unused)}")
    print("wrote main.tex, refs.bib, paper-rendered.md"
          + (", NMAFC-paper.docx" if docx_written else ""))
    if not docx_written:
        print("SKIPPED      NMAFC-paper.docx is open in Word. Close it and "
              "run again to refresh it.")


if __name__ == "__main__":
    main()
